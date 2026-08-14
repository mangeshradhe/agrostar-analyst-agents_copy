#!/usr/bin/env python3
"""
End-to-end pitch pipeline bot for AgroStar new-product training material.

Given an xlsx of new products (Product Name, Pitch Type, Pitch Text,
Source of Pitch, Language, Key Selling Points, Objection Handling — Hindi
row per product), this:

  1. translate   -> Hinglish/Gujarati/Marathi/Telugu/Kannada versions of each
                     product's pitch content (+ a short transliterated display
                     name per language)
  2. merge-xlsx  -> appends the new 6-language rows to pitches_final.xlsx
  3. scripts     -> Hinglish Field-vs-Retailer roleplay training script per product
  4. build-json  -> one JSON file per (product, language) row for the app
  5. thumbnails  -> searches the web and downloads a product photo per product
  6. productlist -> appends new entries to productlist_<lang>.json
  7. consolidate-scripts -> rebuilds roleplay_scripts_all.xlsx from scripts/*.txt
  8. tracker     -> rebuilds audio_thumbnail_tracker.csv/.xlsx from json/*.json

Everything is driven by headless `claude -p` calls (your existing Claude Code
login — no separate ANTHROPIC_API_KEY needed) for the language-dependent
steps, and plain deterministic Python for the reshaping steps.

Usage:
    python3 pipeline.py run <new_products.xlsx> [--workdir DIR] [--jobs N]
    python3 pipeline.py run <new_products.xlsx> --only translate,scripts
    python3 pipeline.py match-audio <drive_downloads_dir> [--workdir DIR]

Run from (or point --workdir at) the folder that already holds
pitches_final.xlsx, scripts/, json/, assets/, and productlist_*.json — e.g.
~/claude_workspace.
"""
import argparse
import csv
import glob
import json
import os
import re
import shutil
import sys
from concurrent.futures import ThreadPoolExecutor, as_completed

import openpyxl
from openpyxl.styles import Alignment, Font

from claude_cli import run_claude, ClaudeCliError
from prompts import translate_prompt, roleplay_prompt, roleplay_prompt_ksp_only, thumbnail_prompt, ALL_LANGS

LANG_CODES = {"English": "en", "Hindi": "hi", "Gujarati": "gu",
              "Marathi": "mr", "Telugu": "te", "Kannada": "kn"}
DEFAULT_SOURCE_LANG = "Hindi"

COLUMNS = ["Product Name", "Pitch Type", "Pitch Text", "Source of Pitch",
           "Language", "Key Selling Points", "Objection Handling"]


def log(msg):
    print(f"[pipeline] {msg}", flush=True)


def slug(name):
    return re.sub(r"[ +:\-]+", "_", name)


# ---------------------------------------------------------------------------
# Stage 0: extract new products from the source xlsx
# ---------------------------------------------------------------------------

def extract_products(xlsx_path):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb.active
    header = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(header)}
    products = []
    seen = set()
    for row in ws.iter_rows(min_row=2, values_only=True):
        name = row[idx["Product Name"]]
        if not name or name in seen:
            continue
        seen.add(name)
        language = row[idx["Language"]] if "Language" in idx else None
        products.append({
            "product": name,
            "language": language or DEFAULT_SOURCE_LANG,
            "pitch_type": row[idx.get("Pitch Type", -1)] if "Pitch Type" in idx else "PRODUCT",
            "pitch_text": row[idx.get("Pitch Text", -1)] if "Pitch Text" in idx else "",
            "source_of_pitch": row[idx.get("Source of Pitch", -1)] if "Source of Pitch" in idx else "",
            "ksp": row[idx["Key Selling Points"]] or "",
            "oh": row[idx["Objection Handling"]] or "",
        })
    unknown = sorted({p["language"] for p in products} - set(ALL_LANGS))
    if unknown:
        raise ValueError(
            f"Source xlsx has products in language(s) not in {ALL_LANGS}: {unknown}. "
            "Add support for that language or fix the Language column."
        )
    return products


# ---------------------------------------------------------------------------
# Stage 1: translate (parallel headless claude -p calls)
# ---------------------------------------------------------------------------

def translate_one(product, out_dir):
    out_path = os.path.join(out_dir, f"{slug(product['product'])}.json")
    if os.path.exists(out_path):
        log(f"translate: {product['product']} already done, skipping")
        return out_path
    prompt = translate_prompt(product["product"], product["language"], product["ksp"], product["oh"], out_path)
    run_claude(prompt, allowed_tools=["Write"], timeout=600)
    if not os.path.exists(out_path):
        raise ClaudeCliError(f"translate: expected output not found for {product['product']}")
    return out_path


def load_rows_by_lang(product, translations_dir):
    """{lang: {"ksp":..., "oh":...}} for all 6 ALL_LANGS, combining the
    product's own source-language content with its translation JSON."""
    tpath = os.path.join(translations_dir, f"{slug(product['product'])}.json")
    if not os.path.exists(tpath):
        return None, None
    with open(tpath, encoding="utf-8") as f:
        t = json.load(f)

    source_lang = product["language"]
    rows_by_lang = {source_lang: {"ksp": product["ksp"], "oh": product["oh"]}}
    for lang in ALL_LANGS:
        if lang == source_lang:
            continue
        rows_by_lang[lang] = {"ksp": t[lang]["ksp"], "oh": t[lang]["oh"]}
    return rows_by_lang, t


def stage_translate(products, out_dir, jobs):
    os.makedirs(out_dir, exist_ok=True)
    results = {}
    with ThreadPoolExecutor(max_workers=jobs) as ex:
        futs = {ex.submit(translate_one, p, out_dir): p for p in products}
        for fut in as_completed(futs):
            p = futs[fut]
            try:
                path = fut.result()
                results[p["product"]] = path
                log(f"translate: done -> {p['product']}")
            except Exception as e:
                log(f"translate: FAILED for {p['product']}: {e}")
    return results


# ---------------------------------------------------------------------------
# Stage 2: merge into the master pitches_final.xlsx (append, never overwrite)
# ---------------------------------------------------------------------------

def stage_merge_xlsx(products, translations_dir, pitches_xlsx):
    if os.path.exists(pitches_xlsx):
        wb = openpyxl.load_workbook(pitches_xlsx)
        ws = wb.active
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Pitches"
        ws.append(COLUMNS)
        for cell in ws[1]:
            cell.font = Font(bold=True)

    existing_names = {row[0] for row in ws.iter_rows(min_row=2, values_only=True)}

    added = 0
    for p in products:
        if p["product"] in existing_names:
            log(f"merge-xlsx: {p['product']} already present, skipping")
            continue

        rows_by_lang, _ = load_rows_by_lang(p, translations_dir)
        if rows_by_lang is None:
            log(f"merge-xlsx: no translation for {p['product']}, skipping")
            continue

        for lang in ALL_LANGS:
            r = rows_by_lang[lang]
            ws.append([
                p["product"], p["pitch_type"], p["pitch_text"], p["source_of_pitch"],
                lang, r["ksp"], r["oh"],
            ])
            added += 1

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    widths = {"A": 18, "B": 12, "C": 14, "D": 14, "E": 12, "F": 45, "G": 55}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

    wb.save(pitches_xlsx)
    log(f"merge-xlsx: appended {added} rows -> {pitches_xlsx}")


# ---------------------------------------------------------------------------
# Stage 3: roleplay scripts
# ---------------------------------------------------------------------------

def script_one(product, scripts_dir):
    out_path = os.path.join(scripts_dir, f"{product['product']}.txt")
    if os.path.exists(out_path):
        log(f"scripts: {product['product']} already done, skipping")
        return out_path
    if (product["oh"] or "").strip():
        prompt = roleplay_prompt(product["product"], product["language"], product["ksp"], product["oh"], out_path)
    else:
        prompt = roleplay_prompt_ksp_only(product["product"], product["language"], product["ksp"], out_path)
    run_claude(prompt, allowed_tools=["Write"], timeout=600)
    if not os.path.exists(out_path):
        raise ClaudeCliError(f"scripts: expected output not found for {product['product']}")
    return out_path


def stage_scripts(products, scripts_dir, jobs):
    os.makedirs(scripts_dir, exist_ok=True)
    with ThreadPoolExecutor(max_workers=jobs) as ex:
        futs = {ex.submit(script_one, p, scripts_dir): p for p in products}
        for fut in as_completed(futs):
            p = futs[fut]
            try:
                fut.result()
                log(f"scripts: done -> {p['product']}")
            except Exception as e:
                log(f"scripts: FAILED for {p['product']}: {e}")


# ---------------------------------------------------------------------------
# Stage 4: per-row JSON files (deterministic)
# ---------------------------------------------------------------------------

def parse_key_points(cell):
    points = []
    for line in cell.split("\n"):
        line = line.strip()
        if not line:
            continue
        line = line.lstrip("-•").strip().strip("\"'").strip()
        if line:
            points.append(line)
    return points


def parse_objections(cell):
    text = cell.strip()
    if len(text) >= 2 and text[0] == '"' and text[-1] == '"':
        text = text[1:-1]
    blocks = re.split(r"⚠️?", text)
    objections = []
    for block in blocks:
        block = block.strip()
        if not block:
            continue
        m = re.search(r"✅|✓", block)
        if not m:
            continue
        objection = block[: m.start()].strip().strip("\"'\n\t ").strip()
        rebuttal = block[m.end():].strip().strip("\"'\n\t ").strip()
        if not objection and not rebuttal:
            continue
        objections.append({"objection": objection, "rebuttal": rebuttal})
    return objections


def stage_build_json(products, translations_dir, json_dir):
    os.makedirs(json_dir, exist_ok=True)
    written = 0
    for p in products:
        rows_by_lang, _ = load_rows_by_lang(p, translations_dir)
        if rows_by_lang is None:
            continue

        uname = slug(p["product"])
        for lang in ALL_LANGS:
            langcode = LANG_CODES[lang]
            r = rows_by_lang[lang]
            obj = {
                "name": p["product"],
                "thumbnail": f"assets/{uname}.jpg",
                "audio": f"audio/{uname}_{langcode}.wav",
                "keyPoints": parse_key_points(r["ksp"]),
                "objections": parse_objections(r["oh"]),
            }
            fname = f"{uname}_{langcode}.json"
            with open(os.path.join(json_dir, fname), "w", encoding="utf-8") as f:
                json.dump(obj, f, ensure_ascii=False, indent=2)
            written += 1
    log(f"build-json: wrote {written} files -> {json_dir}")


# ---------------------------------------------------------------------------
# Stage 5: thumbnails
# ---------------------------------------------------------------------------

def thumbnail_one(product, assets_dir):
    asset_path = os.path.join(assets_dir, f"{slug(product['product'])}.jpg")
    if os.path.exists(asset_path) and os.path.getsize(asset_path) > 5000:
        log(f"thumbnails: {product['product']} already have a real file, skipping")
        return asset_path
    context = (product["ksp"] or "")[:200]
    prompt = thumbnail_prompt(product["product"], f"an agri-input product. Context: {context}", asset_path)
    run_claude(prompt, allowed_tools=["WebSearch", "Bash"], timeout=300)
    if not os.path.exists(asset_path) or os.path.getsize(asset_path) < 5000:
        raise ClaudeCliError(f"thumbnails: expected output not found (or too small) for {product['product']}")
    return asset_path


def stage_thumbnails(products, assets_dir, jobs):
    os.makedirs(assets_dir, exist_ok=True)
    with ThreadPoolExecutor(max_workers=jobs) as ex:
        futs = {ex.submit(thumbnail_one, p, assets_dir): p for p in products}
        for fut in as_completed(futs):
            p = futs[fut]
            try:
                fut.result()
                log(f"thumbnails: done -> {p['product']}")
            except Exception as e:
                log(f"thumbnails: FAILED for {p['product']}: {e}")


# ---------------------------------------------------------------------------
# Stage 6: productlist_<lang>.json
# ---------------------------------------------------------------------------

def stage_productlist(products, translations_dir, workdir, assets_dir):
    asset_files = set(os.listdir(assets_dir)) if os.path.isdir(assets_dir) else set()
    missing = []

    for lang, code in LANG_CODES.items():
        path = os.path.join(workdir, f"productlist_{code}.json")
        if not os.path.exists(path):
            log(f"productlist: {path} not found, skipping")
            continue
        with open(path, encoding="utf-8") as f:
            data = json.load(f)
        existing_names = {e["name"] for e in data}
        old_count = len(data)

        for p in products:
            if p["product"] in existing_names:
                continue
            uname = slug(p["product"])
            fname = f"{uname}.jpg"
            if fname not in asset_files:
                missing.append(fname)

            if code == "en":
                display = p["product"]
            else:
                tpath = os.path.join(translations_dir, f"{uname}.json")
                display = p["product"]
                if os.path.exists(tpath):
                    with open(tpath, encoding="utf-8") as tf:
                        t = json.load(tf)
                    display = t.get(lang, {}).get("displayName", p["product"])

            data.append({
                "name": p["product"],
                "displayName": display,
                "thumbnail": f"assets/{uname}.jpg",
                "category": "New Products",
                "tag": None,
            })

        with open(path, "w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False, indent=2)
        log(f"productlist_{code}.json: {old_count} -> {len(data)}")

    if missing:
        log("WARNING: missing thumbnail files referenced by productlist entries:")
        for m in sorted(set(missing)):
            log(f"  - {m}")


# ---------------------------------------------------------------------------
# Stage 7: consolidate all roleplay scripts into one xlsx (full rebuild)
# ---------------------------------------------------------------------------

def stage_consolidate_scripts(workdir):
    scripts_dir = os.path.join(workdir, "scripts")
    out_path = os.path.join(workdir, "roleplay_scripts_all.xlsx")
    if not os.path.isdir(scripts_dir):
        log(f"consolidate-scripts: {scripts_dir} not found, skipping")
        return

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Roleplay Scripts"
    ws.append(["Product Name", "Roleplay Script"])
    for cell in ws[1]:
        cell.font = Font(bold=True)

    files = sorted(glob.glob(os.path.join(scripts_dir, "*.txt")))
    for f in files:
        product = os.path.splitext(os.path.basename(f))[0]
        with open(f, encoding="utf-8") as fh:
            content = fh.read()
        ws.append([product, content])

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 100
    for i in range(2, ws.max_row + 1):
        ws.row_dimensions[i].height = 400

    wb.save(out_path)
    log(f"consolidate-scripts: wrote {len(files)} products -> {out_path}")


# ---------------------------------------------------------------------------
# Stage 8: audio/thumbnail tracker (full rebuild from json/)
# ---------------------------------------------------------------------------

def stage_tracker(workdir):
    json_dir = os.path.join(workdir, "json")
    if not os.path.isdir(json_dir):
        log(f"tracker: {json_dir} not found, skipping")
        return

    lang_order = ["English", "Hindi", "Gujarati", "Marathi", "Telugu", "Kannada"]
    lang_names_by_code = {v: k for k, v in LANG_CODES.items()}

    rows = []
    for fp in sorted(glob.glob(os.path.join(json_dir, "*.json"))):
        with open(fp, encoding="utf-8") as f:
            d = json.load(f)
        fname = os.path.basename(fp)
        langcode = fname.rsplit("_", 1)[1].replace(".json", "")
        rows.append({
            "File Name": fname,
            "Name": d["name"],
            "Audio": d["audio"],
            "Language": lang_names_by_code[langcode],
            "Thumbnail": d["thumbnail"],
        })

    rows.sort(key=lambda r: (r["Name"], lang_order.index(r["Language"])))
    fields = ["File Name", "Name", "Audio", "Language", "Thumbnail"]

    csv_path = os.path.join(workdir, "audio_thumbnail_tracker.csv")
    with open(csv_path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=fields)
        w.writeheader()
        w.writerows(rows)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Tracker"
    ws.append(fields)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for r in rows:
        ws.append([r[k] for k in fields])
    widths = {"A": 30, "B": 20, "C": 35, "D": 12, "E": 30}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    wb.save(os.path.join(workdir, "audio_thumbnail_tracker.xlsx"))

    log(f"tracker: wrote {len(rows)} rows -> audio_thumbnail_tracker.csv/.xlsx")


# ---------------------------------------------------------------------------
# Audio matching helper (semi-manual, since Drive filenames aren't predictable)
# ---------------------------------------------------------------------------

def match_audio(drive_dir, workdir):
    json_dir = os.path.join(workdir, "json")
    audio_dir = os.path.join(workdir, "audio")
    os.makedirs(audio_dir, exist_ok=True)

    expected = []
    for fname in sorted(os.listdir(json_dir)):
        if not fname.endswith(".json"):
            continue
        with open(os.path.join(json_dir, fname), encoding="utf-8") as f:
            d = json.load(f)
        audio_name = os.path.basename(d["audio"])
        expected.append((d["name"], audio_name))

    drive_files = sorted(
        f for f in os.listdir(drive_dir)
        if os.path.isfile(os.path.join(drive_dir, f))
    )

    print(f"\n{len(expected)} expected audio files, {len(drive_files)} files in {drive_dir}\n")
    try:
        for product_name, audio_name in expected:
            target_path = os.path.join(audio_dir, audio_name)
            if os.path.exists(target_path):
                continue

            # Best-effort suggestion: files whose name contains a token from the product name.
            tokens = [t.lower() for t in re.split(r"[\s_+:-]+", product_name) if len(t) > 2]
            candidates = [f for f in drive_files if any(t in f.lower() for t in tokens)]

            print(f"Need: {audio_name}  (product: {product_name})")
            if candidates:
                for i, c in enumerate(candidates, 1):
                    print(f"  [{i}] {c}")
            else:
                print("  (no filename match found — list all files with 'l')")

            choice = input("  Pick number, 'l' to list all files, or Enter to skip: ").strip()
            if choice == "":
                continue
            if choice.lower() == "l":
                for i, f in enumerate(drive_files, 1):
                    print(f"  [{i}] {f}")
                choice = input("  Pick number, or Enter to skip: ").strip()
                if choice == "":
                    continue
                src = drive_files[int(choice) - 1]
            else:
                src = candidates[int(choice) - 1]

            shutil.copyfile(os.path.join(drive_dir, src), target_path)
            print(f"  -> copied and renamed to audio/{audio_name}\n")
    except (EOFError, KeyboardInterrupt):
        print("\n\nStopped early — progress so far is saved in audio/. Run the same command again to continue.")
        return

    print(f"\nDone. Renamed files are staged in {audio_dir}/")
    print("To upload to S3 once you have credentials:")
    print(f"  aws s3 sync {audio_dir} s3://<your-bucket>/audio/")


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

STAGE_ORDER = ["translate", "merge-xlsx", "scripts", "build-json", "thumbnails", "productlist",
               "consolidate-scripts", "tracker"]


def run(args):
    workdir = os.path.abspath(args.workdir)
    products = extract_products(args.xlsx)
    log(f"found {len(products)} unique products in {args.xlsx}")

    only = set(args.only.split(",")) if args.only else set(STAGE_ORDER)

    translations_dir = os.path.join(workdir, "pipeline_out", "translations")
    scripts_dir = os.path.join(workdir, "scripts")
    json_dir = os.path.join(workdir, "json")
    assets_dir = os.path.join(workdir, "assets")
    pitches_xlsx = os.path.join(workdir, args.pitches_xlsx)

    if "translate" in only:
        stage_translate(products, translations_dir, args.jobs)
    if "merge-xlsx" in only:
        stage_merge_xlsx(products, translations_dir, pitches_xlsx)
    if "scripts" in only:
        stage_scripts(products, scripts_dir, args.jobs)
    if "build-json" in only:
        stage_build_json(products, translations_dir, json_dir)
    if "thumbnails" in only:
        stage_thumbnails(products, assets_dir, args.jobs)
    if "productlist" in only:
        stage_productlist(products, translations_dir, workdir, assets_dir)
    if "consolidate-scripts" in only:
        stage_consolidate_scripts(workdir)
    if "tracker" in only:
        stage_tracker(workdir)

    log("pipeline run complete.")


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = parser.add_subparsers(dest="cmd", required=True)

    p_run = sub.add_parser("run", help="Run the full pipeline for a new-products xlsx")
    p_run.add_argument("xlsx", help="Path to the new-products xlsx from the training team")
    p_run.add_argument("--workdir", default=".", help="Folder holding pitches_final.xlsx, scripts/, json/, assets/, productlist_*.json (default: current dir)")
    p_run.add_argument("--pitches-xlsx", default="pitches_final.xlsx", help="Master pitches xlsx filename inside workdir")
    p_run.add_argument("--jobs", type=int, default=6, help="Parallel claude -p calls (default 6)")
    p_run.add_argument("--only", default=None, help=f"Comma-separated subset of stages to run: {','.join(STAGE_ORDER)}")

    p_audio = sub.add_parser("match-audio", help="Interactively match/rename Drive audio downloads into audio/")
    p_audio.add_argument("drive_dir", help="Folder where you downloaded the training team's audio files from Drive")
    p_audio.add_argument("--workdir", default=".", help="Folder holding json/ (default: current dir)")

    args = parser.parse_args()

    if args.cmd == "run":
        run(args)
    elif args.cmd == "match-audio":
        match_audio(args.drive_dir, os.path.abspath(args.workdir))


if __name__ == "__main__":
    main()
